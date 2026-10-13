import os
from datetime import datetime, date, timedelta
from functools import wraps
from io import BytesIO

from flask import Flask, render_template, request, redirect, url_for, flash, session, send_file
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
from sqlalchemy import func

app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'change-this-secret-key-before-production')
# SQLite database. Set DB_PATH to choose a persistent location when hosting online.
# Local default: instance/supermart.db (created automatically).
instance_dir = os.path.join(app.root_path, 'instance')
os.makedirs(instance_dir, exist_ok=True)
database_path = os.path.abspath(os.environ.get('DB_PATH', os.path.join(instance_dir, 'supermart.db')))
os.makedirs(os.path.dirname(database_path), exist_ok=True)
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///' + database_path
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SQLALCHEMY_ENGINE_OPTIONS'] = {'connect_args': {'check_same_thread': False}}
db = SQLAlchemy(app)

class Supermarket(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(150), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    supermarket_id = db.Column(db.Integer, db.ForeignKey('supermarket.id'), nullable=False, index=True)
    name = db.Column(db.String(100), nullable=False)
    email = db.Column(db.String(160), unique=True, nullable=False, index=True)
    password_hash = db.Column(db.String(255), nullable=False)

class Product(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    supermarket_id = db.Column(db.Integer, db.ForeignKey('supermarket.id'), nullable=False, index=True)
    name = db.Column(db.String(150), nullable=False)
    sku = db.Column(db.String(60), nullable=True)
    unit = db.Column(db.String(20), default='pcs')
    cost_price = db.Column(db.Float, default=0)
    selling_price = db.Column(db.Float, default=0)
    stock_qty = db.Column(db.Float, default=0)
    low_stock_level = db.Column(db.Float, default=5)
    active = db.Column(db.Boolean, default=True)
    __table_args__ = (db.UniqueConstraint('supermarket_id', 'name', name='uq_product_name_store'),)

class Purchase(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    supermarket_id = db.Column(db.Integer, db.ForeignKey('supermarket.id'), nullable=False, index=True)
    product_id = db.Column(db.Integer, db.ForeignKey('product.id'), nullable=False)
    supplier = db.Column(db.String(150), default='')
    quantity = db.Column(db.Float, nullable=False)
    unit_cost = db.Column(db.Float, nullable=False)
    total = db.Column(db.Float, nullable=False)
    purchase_date = db.Column(db.Date, default=date.today, index=True)
    note = db.Column(db.String(250), default='')
    paid = db.Column(db.Float, default=0, nullable=False)
    payment_method = db.Column(db.String(30), default='Credit', nullable=False)
    product = db.relationship('Product')

class Sale(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    supermarket_id = db.Column(db.Integer, db.ForeignKey('supermarket.id'), nullable=False, index=True)
    invoice_no = db.Column(db.String(40), nullable=False)
    sale_date = db.Column(db.DateTime, default=datetime.now, index=True)
    customer_name = db.Column(db.String(150), default='Walk-in customer')
    payment_method = db.Column(db.String(30), default='Cash')
    subtotal = db.Column(db.Float, nullable=False)
    discount = db.Column(db.Float, default=0)
    total = db.Column(db.Float, nullable=False)
    paid = db.Column(db.Float, default=0)
    items = db.relationship('SaleItem', backref='sale', cascade='all, delete-orphan')
    __table_args__ = (db.UniqueConstraint('supermarket_id', 'invoice_no', name='uq_invoice_store'),)

class SaleItem(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    sale_id = db.Column(db.Integer, db.ForeignKey('sale.id'), nullable=False)
    product_id = db.Column(db.Integer, db.ForeignKey('product.id'), nullable=False)
    product_name = db.Column(db.String(150), nullable=False)
    quantity = db.Column(db.Float, nullable=False)
    unit_price = db.Column(db.Float, nullable=False)
    cost_price = db.Column(db.Float, nullable=False)
    line_total = db.Column(db.Float, nullable=False)
    product = db.relationship('Product')

class Expense(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    supermarket_id = db.Column(db.Integer, db.ForeignKey('supermarket.id'), nullable=False, index=True)
    expense_date = db.Column(db.Date, default=date.today, index=True)
    category = db.Column(db.String(100), nullable=False)
    description = db.Column(db.String(250), default='')
    amount = db.Column(db.Float, nullable=False)
    payment_method = db.Column(db.String(30), default='Cash')

with app.app_context():
    db.create_all()
    # Lightweight SQLite schema migration for existing installations.
    from sqlalchemy import inspect
    inspector = inspect(db.engine)
    purchase_columns = {column['name'] for column in inspector.get_columns('purchase')}
    with db.engine.begin() as connection:
        if 'paid' not in purchase_columns:
            connection.exec_driver_sql('ALTER TABLE purchase ADD COLUMN paid FLOAT NOT NULL DEFAULT 0')
        if 'payment_method' not in purchase_columns:
            connection.exec_driver_sql("ALTER TABLE purchase ADD COLUMN payment_method VARCHAR(30) NOT NULL DEFAULT 'Credit'")

def store_id():
    return session.get('supermarket_id')

def login_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not session.get('user_id'):
            return redirect(url_for('login'))
        return fn(*args, **kwargs)
    return wrapped

def money(value):
    return f"₹{(value or 0):,.2f}"
app.jinja_env.filters['money'] = money

def date_range():
    today = date.today()
    start_raw = request.args.get('start', '')
    end_raw = request.args.get('end', '')
    try: start = datetime.strptime(start_raw, '%Y-%m-%d').date() if start_raw else today.replace(day=1)
    except ValueError: start = today.replace(day=1)
    try: end = datetime.strptime(end_raw, '%Y-%m-%d').date() if end_raw else today
    except ValueError: end = today
    if end < start: start, end = end, start
    return start, end

@app.context_processor
def common_context():
    return {'store_name': session.get('store_name', 'SuperMart'), 'today': date.today().isoformat()}

@app.route('/')
def index():
    if not session.get('user_id'): return redirect(url_for('login'))
    return redirect(url_for('dashboard'))

@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        store = request.form.get('store_name','').strip()
        name = request.form.get('name','').strip()
        email = request.form.get('email','').strip().lower()
        password = request.form.get('password','')
        if not all([store, name, email, password]):
            flash('Please complete every field.', 'error')
        elif len(password) < 8:
            flash('Password must contain at least 8 characters.', 'error')
        elif User.query.filter_by(email=email).first():
            flash('This email is already registered. Please sign in.', 'error')
        else:
            supermarket = Supermarket(name=store)
            db.session.add(supermarket); db.session.flush()
            user = User(supermarket_id=supermarket.id, name=name, email=email, password_hash=generate_password_hash(password))
            db.session.add(user); db.session.commit()
            session.update(user_id=user.id, supermarket_id=supermarket.id, store_name=supermarket.name, user_name=user.name)
            flash('Your supermarket workspace is ready.', 'success')
            return redirect(url_for('dashboard'))
    return render_template('register.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        email = request.form.get('email','').strip().lower()
        password = request.form.get('password','')
        user = User.query.filter_by(email=email).first()
        if user and check_password_hash(user.password_hash, password):
            supermarket = db.session.get(Supermarket, user.supermarket_id)
            session.clear(); session.update(user_id=user.id, supermarket_id=user.supermarket_id, store_name=supermarket.name, user_name=user.name)
            return redirect(url_for('dashboard'))
        flash('Incorrect email or password.', 'error')
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear(); return redirect(url_for('login'))

@app.route('/dashboard')
@login_required
def dashboard():
    sid = store_id(); today = date.today()
    sales_today = db.session.query(func.coalesce(func.sum(Sale.total),0)).filter(Sale.supermarket_id==sid, func.date(Sale.sale_date)==today.isoformat()).scalar() or 0
    purchases_today = db.session.query(func.coalesce(func.sum(Purchase.total),0)).filter_by(supermarket_id=sid, purchase_date=today).scalar() or 0
    expenses_today = db.session.query(func.coalesce(func.sum(Expense.amount),0)).filter_by(supermarket_id=sid, expense_date=today).scalar() or 0
    product_count = Product.query.filter_by(supermarket_id=sid, active=True).count()
    low_stock = Product.query.filter(Product.supermarket_id==sid, Product.active.is_(True), Product.stock_qty <= Product.low_stock_level).order_by(Product.stock_qty.asc()).limit(8).all()
    recent_sales = Sale.query.filter_by(supermarket_id=sid).order_by(Sale.sale_date.desc()).limit(6).all()
    return render_template('dashboard.html', sales_today=sales_today, purchases_today=purchases_today, expenses_today=expenses_today, product_count=product_count, low_stock=low_stock, recent_sales=recent_sales)

@app.route('/products', methods=['GET','POST'])
@login_required
def products():
    sid=store_id()
    if request.method=='POST':
        name=request.form.get('name','').strip(); sku=request.form.get('sku','').strip()
        try:
            cost=float(request.form.get('cost_price',0)); price=float(request.form.get('selling_price',0)); qty=float(request.form.get('stock_qty',0)); low=float(request.form.get('low_stock_level',5))
            if not name or min(cost,price,qty,low)<0: raise ValueError()
        except ValueError:
            flash('Enter a product name and valid non-negative numbers.', 'error'); return redirect(url_for('products'))
        existing=Product.query.filter_by(supermarket_id=sid,name=name).first()
        if existing:
            flash('A product with this name already exists.', 'error')
        else:
            db.session.add(Product(supermarket_id=sid,name=name,sku=sku or None,unit=request.form.get('unit','pcs'),cost_price=cost,selling_price=price,stock_qty=qty,low_stock_level=low)); db.session.commit(); flash('Product added.', 'success')
        return redirect(url_for('products'))
    q=request.args.get('q','').strip()
    query=Product.query.filter_by(supermarket_id=sid,active=True)
    if q: query=query.filter(Product.name.ilike(f'%{q}%'))
    return render_template('products.html', products=query.order_by(Product.name).all(), q=q)

@app.route('/purchases', methods=['GET','POST'])
@login_required
def purchases():
    sid=store_id()
    if request.method=='POST':
        try:
            product=Product.query.filter_by(id=int(request.form['product_id']),supermarket_id=sid,active=True).first_or_404()
            qty=float(request.form['quantity']); cost=float(request.form['unit_cost'])
            if qty<=0 or cost<0: raise ValueError()
            pdate=datetime.strptime(request.form.get('purchase_date') or date.today().isoformat(),'%Y-%m-%d').date()
            total=qty*cost
            db.session.add(Purchase(supermarket_id=sid,product_id=product.id,supplier=request.form.get('supplier','').strip(),quantity=qty,unit_cost=cost,total=total,purchase_date=pdate,note=request.form.get('note','').strip(),paid=min(max(float(request.form.get('paid',0)),0),total),payment_method=request.form.get('payment_method','Credit')))
            product.stock_qty += qty; product.cost_price=cost
            db.session.commit(); flash('Purchase saved and stock updated.', 'success')
        except (ValueError, KeyError):
            db.session.rollback(); flash('Please enter valid purchase details.', 'error')
        return redirect(url_for('purchases'))
    items=Purchase.query.filter_by(supermarket_id=sid).order_by(Purchase.purchase_date.desc(),Purchase.id.desc()).limit(100).all()
    return render_template('purchases.html', products=Product.query.filter_by(supermarket_id=sid,active=True).order_by(Product.name).all(), purchases=items)

@app.route('/sales', methods=['GET','POST'])
@login_required
def sales():
    sid=store_id()
    products_list=Product.query.filter_by(supermarket_id=sid,active=True).order_by(Product.name).all()
    if request.method=='POST':
        try:
            ids=request.form.getlist('product_id[]'); qtys=request.form.getlist('quantity[]')
            discount=float(request.form.get('discount',0)); paid=float(request.form.get('paid',0))
            if discount<0 or paid<0: raise ValueError()
            line_data=[]; subtotal=0; requested_by_product={}
            for pid, qraw in zip(ids,qtys):
                if not pid or not qraw: continue
                product=Product.query.filter_by(id=int(pid),supermarket_id=sid,active=True).first()
                qty=float(qraw)
                if not product or qty<=0: raise ValueError('Invalid product or quantity.')
                requested_by_product[product.id] = requested_by_product.get(product.id, 0) + qty
                if product.stock_qty < requested_by_product[product.id]: raise ValueError(f'Not enough stock for {product.name}. Available: {product.stock_qty:g}')
                line_total=qty*product.selling_price; subtotal+=line_total
                line_data.append((product,qty,line_total))
            if not line_data: raise ValueError('Add at least one product.')
            if discount>subtotal: raise ValueError('Discount cannot exceed subtotal.')
            total=subtotal-discount
            invoice='SM'+datetime.now().strftime('%y%m%d%H%M%S')+str(session['user_id'])
            sale=Sale(supermarket_id=sid,invoice_no=invoice,customer_name=request.form.get('customer_name','').strip() or 'Walk-in customer',payment_method=request.form.get('payment_method','Cash'),subtotal=subtotal,discount=discount,total=total,paid=min(paid,total))
            db.session.add(sale); db.session.flush()
            for product,qty,line_total in line_data:
                db.session.add(SaleItem(sale_id=sale.id,product_id=product.id,product_name=product.name,quantity=qty,unit_price=product.selling_price,cost_price=product.cost_price,line_total=line_total))
                product.stock_qty-=qty
            db.session.commit(); return redirect(url_for('sale_receipt',sale_id=sale.id))
        except (ValueError, KeyError) as exc:
            db.session.rollback(); flash(str(exc) or 'Please enter valid sale details.', 'error')
            return redirect(url_for('sales'))
    return render_template('sales.html', products=products_list)

@app.route('/sales/<int:sale_id>/receipt')
@login_required
def sale_receipt(sale_id):
    sale=Sale.query.filter_by(id=sale_id,supermarket_id=store_id()).first_or_404()
    return render_template('receipt.html', sale=sale)

@app.route('/sales/history')
@login_required
def sales_history():
    start,end=date_range(); sid=store_id()
    rows=Sale.query.filter(Sale.supermarket_id==sid,func.date(Sale.sale_date)>=start.isoformat(),func.date(Sale.sale_date)<=end.isoformat()).order_by(Sale.sale_date.desc()).all()
    return render_template('sales_history.html', sales=rows,start=start.isoformat(),end=end.isoformat(),total=sum(x.total for x in rows))

@app.route('/expenses', methods=['GET','POST'])
@login_required
def expenses():
    sid=store_id()
    if request.method=='POST':
        try:
            amount=float(request.form.get('amount',0)); edate=datetime.strptime(request.form.get('expense_date') or date.today().isoformat(),'%Y-%m-%d').date()
            category=request.form.get('category','').strip()
            if amount<=0 or not category: raise ValueError()
            db.session.add(Expense(supermarket_id=sid,expense_date=edate,category=category,description=request.form.get('description','').strip(),amount=amount,payment_method=request.form.get('payment_method','Cash'))); db.session.commit(); flash('Expense saved.', 'success')
        except ValueError:
            db.session.rollback(); flash('Enter a category and a valid amount.', 'error')
        return redirect(url_for('expenses'))
    rows=Expense.query.filter_by(supermarket_id=sid).order_by(Expense.expense_date.desc(),Expense.id.desc()).limit(100).all()
    return render_template('expenses.html',expenses=rows)

@app.route('/reports')
@login_required
def reports():
    start,end=date_range(); sid=store_id()
    sales_rows=Sale.query.filter(Sale.supermarket_id==sid,func.date(Sale.sale_date)>=start.isoformat(),func.date(Sale.sale_date)<=end.isoformat()).order_by(Sale.sale_date.desc()).all()
    purchase_rows=Purchase.query.filter(Purchase.supermarket_id==sid,Purchase.purchase_date>=start,Purchase.purchase_date<=end).order_by(Purchase.purchase_date.desc()).all()
    expense_rows=Expense.query.filter(Expense.supermarket_id==sid,Expense.expense_date>=start,Expense.expense_date<=end).order_by(Expense.expense_date.desc()).all()
    cogs=sum(item.quantity*item.cost_price for sale in sales_rows for item in sale.items)
    sales_total=sum(s.total for s in sales_rows); purchase_total=sum(p.total for p in purchase_rows); expense_total=sum(e.amount for e in expense_rows)
    return render_template('reports.html',start=start.isoformat(),end=end.isoformat(),sales=sales_rows,purchases=purchase_rows,expenses=expense_rows,sales_total=sales_total,purchase_total=purchase_total,expense_total=expense_total,cogs=cogs,gross_profit=sales_total-cogs,net_estimate=sales_total-cogs-expense_total)

@app.route('/reports/export.xlsx')
@login_required
def export_reports():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    start,end=date_range(); sid=store_id(); wb=Workbook(); ws=wb.active; ws.title='Summary'
    ws.append([session.get('store_name','SuperMart'),'Report summary']); ws.append(['Period',f'{start} to {end}'])
    sales_rows=Sale.query.filter(Sale.supermarket_id==sid,func.date(Sale.sale_date)>=start.isoformat(),func.date(Sale.sale_date)<=end.isoformat()).all()
    purchase_rows=Purchase.query.filter(Purchase.supermarket_id==sid,Purchase.purchase_date>=start,Purchase.purchase_date<=end).all()
    expense_rows=Expense.query.filter(Expense.supermarket_id==sid,Expense.expense_date>=start,Expense.expense_date<=end).all()
    sales_total=sum(s.total for s in sales_rows); purchase_total=sum(p.total for p in purchase_rows); expense_total=sum(e.amount for e in expense_rows)
    cogs=sum(i.quantity*i.cost_price for s in sales_rows for i in s.items)
    for row in [['Sales',sales_total],['Purchases',purchase_total],['Other expenses',expense_total],['Cost of goods sold (estimated)',cogs],['Gross profit (estimated)',sales_total-cogs],['Net estimate after expenses',sales_total-cogs-expense_total]]: ws.append(row)
    ws2=wb.create_sheet('Sales'); ws2.append(['Invoice','Date','Customer','Payment','Subtotal','Discount','Total','Paid'])
    for s in sales_rows: ws2.append([s.invoice_no,s.sale_date.strftime('%Y-%m-%d %H:%M'),s.customer_name,s.payment_method,s.subtotal,s.discount,s.total,s.paid])
    ws3=wb.create_sheet('Purchases'); ws3.append(['Date','Product','Supplier','Quantity','Unit cost','Total','Paid','Outstanding','Payment method','Note'])
    for p in purchase_rows: ws3.append([p.purchase_date.isoformat(),p.product.name,p.supplier,p.quantity,p.unit_cost,p.total,p.paid or 0,max(0,p.total-(p.paid or 0)),p.payment_method,p.note])
    ws4=wb.create_sheet('Expenses'); ws4.append(['Date','Category','Description','Payment','Amount'])
    for e in expense_rows: ws4.append([e.expense_date.isoformat(),e.category,e.description,e.payment_method,e.amount])
    for sheet in wb.worksheets:
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='203B53')
        for col in sheet.columns:
            width=min(max(max(len(str(c.value or '')) for c in col)+2,12),32)
            sheet.column_dimensions[get_column_letter(col[0].column)].width=width
    output=BytesIO(); wb.save(output); output.seek(0)
    return send_file(output,as_attachment=True,download_name=f'supermart_report_{start}_{end}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/accounting')
@login_required
def accounting():
    start,end=date_range(); sid=store_id()
    ss=Sale.query.filter(Sale.supermarket_id==sid,func.date(Sale.sale_date)>=start.isoformat(),func.date(Sale.sale_date)<=end.isoformat()).all()
    ps=Purchase.query.filter(Purchase.supermarket_id==sid,Purchase.purchase_date>=start,Purchase.purchase_date<=end).all()
    es=Expense.query.filter(Expense.supermarket_id==sid,Expense.expense_date>=start,Expense.expense_date<=end).all()
    sales_total=sum(s.total for s in ss); cogs=sum(i.quantity*i.cost_price for s in ss for i in s.items); purchase_total=sum(p.total for p in ps); expense_total=sum(e.amount for e in es)
    return render_template('accounting.html',start=start.isoformat(),end=end.isoformat(),sales_total=sales_total,purchase_total=purchase_total,expense_total=expense_total,cogs=cogs,gross_profit=sales_total-cogs,net_estimate=sales_total-cogs-expense_total)


@app.route('/receivables-payables')
@login_required
def receivables_payables():
    sid = store_id()
    start, end = date_range()
    sales_rows = Sale.query.filter(Sale.supermarket_id == sid, func.date(Sale.sale_date) >= start.isoformat(), func.date(Sale.sale_date) <= end.isoformat()).order_by(Sale.sale_date.desc()).all()
    purchase_rows = Purchase.query.filter(Purchase.supermarket_id == sid, Purchase.purchase_date >= start, Purchase.purchase_date <= end).order_by(Purchase.purchase_date.desc()).all()
    receivables = [sale for sale in sales_rows if sale.total - (sale.paid or 0) > 0.009]
    payables = [purchase for purchase in purchase_rows if purchase.total - (purchase.paid or 0) > 0.009]
    return render_template('receivables_payables.html', start=start.isoformat(), end=end.isoformat(), receivables=receivables, payables=payables,
        receivable_total=sum(s.total - (s.paid or 0) for s in receivables), payable_total=sum(p.total - (p.paid or 0) for p in payables))

@app.route('/receivables-payables/export.xlsx')
@login_required
def export_receivables_payables():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    sid = store_id(); start, end = date_range(); wb = Workbook()
    ar = wb.active; ar.title = 'Accounts Receivable'
    ar.append(['Invoice', 'Date', 'Customer', 'Payment method', 'Invoice total', 'Amount received', 'Outstanding'])
    sales_rows = Sale.query.filter(Sale.supermarket_id == sid, func.date(Sale.sale_date) >= start.isoformat(), func.date(Sale.sale_date) <= end.isoformat()).order_by(Sale.sale_date.desc()).all()
    for sale in sales_rows:
        due = round(sale.total - (sale.paid or 0), 2)
        if due > 0.009: ar.append([sale.invoice_no, sale.sale_date.strftime('%Y-%m-%d %H:%M'), sale.customer_name, sale.payment_method, sale.total, sale.paid or 0, due])
    ap = wb.create_sheet('Accounts Payable')
    ap.append(['Date', 'Supplier', 'Product', 'Payment method', 'Purchase total', 'Amount paid', 'Outstanding'])
    purchase_rows = Purchase.query.filter(Purchase.supermarket_id == sid, Purchase.purchase_date >= start, Purchase.purchase_date <= end).order_by(Purchase.purchase_date.desc()).all()
    for purchase in purchase_rows:
        due = round(purchase.total - (purchase.paid or 0), 2)
        if due > 0.009: ap.append([purchase.purchase_date.isoformat(), purchase.supplier, purchase.product.name, purchase.payment_method, purchase.total, purchase.paid or 0, due])
    for sheet in wb.worksheets:
        sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]: cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='203B53')
        for col in sheet.columns:
            width = min(max(max(len(str(c.value or '')) for c in col) + 2, 12), 32)
            sheet.column_dimensions[get_column_letter(col[0].column)].width = width
    output = BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, as_attachment=True, download_name=f'supermart_ar_ap_{start}_{end}.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/health')
def health(): return {'status':'ok'}

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 5000)), debug=os.environ.get('FLASK_DEBUG')=='1')
